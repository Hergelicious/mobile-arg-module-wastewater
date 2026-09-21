import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
s1=pd.read_csv('S1_tmp.csv')
s2=pd.read_csv('S2_tmp.csv')
assert len(s1)==114
F='Arial'; HF=Font(name=F,bold=True,color='FFFFFF',size=10); BF=Font(name=F,size=10); TF=Font(name=F,bold=True,size=12)
HFILL=PatternFill('solid',fgColor='00224E')   # cividis dark end, matches figures
thin=Side(style='thin',color='BFBFBF')
wb=Workbook()
rd=wb.active; rd.title='README'
desc=[('Supplementary Tables for: Mobile antibiotic resistance genes form a reproducible, integron-associated module in wastewater',None),
('',None),
('Source of input data','Source Data workbook of Zhu et al. (2025) Nat. Commun. 16, 4006 (doi:10.1038/s41467-025-59019-3). Values below are outputs of the analysis code supplied with the manuscript (canonical.py, module.py, controls.py, decode_final.py, rf_fold.py); they are computed statistics, not editable inputs.'),
('Table S1','Per-gene results for the 114 core ARG subtypes (detected in >50% of samples).'),
('Table S2','Summary of all control, robustness and label-free module analyses.'),
('',None),
('Table S1 columns',None),
('gene','ARG subtype name as given in the source data'),
('drug_class / mechanism','Recovered from the source workbook by non-negative least squares (class and mechanism totals are exact sums of subtypes)'),
('prevalence','Fraction of the 226 samples in which the gene was detected'),
('mean_log10_abundance','log10 of the mean per-cell abundance across samples'),
('plant_level_SD','SD across the 142 plants of plant-mean log10 per-cell abundance'),
('mobile_associated','TRUE if the gene belongs to a family classically associated with integrons, plasmids or transposons (curated; see Methods 2.3)'),
('MGE_coupling_index','Within-country partial Spearman correlation with ARG-proximal MGE abundance, controlling for total ARG load'),
('MAG_mobility_index','Fraction of multi-species ARG ORFs in MAGs for the gene\'s class x mechanism (blank where fewer than 10 ORFs)'),
('country_R2','Marginal R2 of country for plant-level log10 abundance'),
('excess_country_R2','country_R2 minus its mean under 9,999 city-block permutations'),
('permutation_P','Per-gene permutation P for country_R2 (9,999 city-block permutations)'),
('SS_share_*','Sequential sum-of-squares shares (country / city within country / plant within city) in the 7 multi-city countries'),
('cluster / in_label_free_module','Label-free clustering of within-country correlations (average linkage, k chosen by silhouette); module = cluster most enriched for mobile-associated genes'),
('rho_intI1_within_country','Within-country Spearman correlation (plant level) with the class 1 integron integrase gene intI1 (pre-specified test B)'),
('',None),
('Table S2 columns','analysis = description; estimate = statistic; P = permutation or test P where applicable ("<0.001" when no permutation exceeded the observed value); note = settings')]
for i,(a,b) in enumerate(desc,1):
    rd.cell(i,1,a).font=TF if i==1 else Font(name=F,bold=bool(b is None and a),size=10)
    if b: rd.cell(i,2,b).font=BF; rd.cell(i,2).alignment=Alignment(wrap_text=True,vertical='top')
    rd.cell(i,1).alignment=Alignment(wrap_text=(i==1),vertical='top')
rd.merge_cells('A1:B1'); rd.row_dimensions[1].height=32
rd.column_dimensions['A'].width=34; rd.column_dimensions['B'].width=110
def write(ws,df,widths,numfmt):
    for j,c in enumerate(df.columns,1):
        cell=ws.cell(1,j,c); cell.font=HF; cell.fill=HFILL; cell.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
    for i,row in enumerate(df.itertuples(index=False),2):
        for j,v in enumerate(row,1):
            if isinstance(v,float) and pd.isna(v): v=None
            cell=ws.cell(i,j,v); cell.font=BF; cell.border=Border(bottom=thin)
            fmt=numfmt.get(df.columns[j-1])
            if fmt and isinstance(v,(int,float)) and not isinstance(v,bool): cell.number_format=fmt
    for j,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(j)].width=w
    ws.freeze_panes='B2'; ws.row_dimensions[1].height=42; ws.auto_filter.ref=ws.dimensions
t1=wb.create_sheet('Table S1')
fm={c:'0.000' for c in s1.columns if s1[c].dtype.kind=='f'}; fm['permutation_P']='0.0000'
write(t1,s1,[30,16,26,11,13,12,11,13,13,11,12,12,12,12,12,9,12,13],fm)
t2=wb.create_sheet('Table S2')
s2o=s2.copy()
write(t2,s2o,[66,22,12,48],{})
for r in range(2,len(s2o)+2): t2.cell(r,1).alignment=Alignment(wrap_text=True); t2.cell(r,4).alignment=Alignment(wrap_text=True)
wb.save('Supplementary_Tables_S1_S2.xlsx')
# verify round-trip
chk1=pd.read_excel('Supplementary_Tables_S1_S2.xlsx',sheet_name='Table S1')
chk2=pd.read_excel('Supplementary_Tables_S1_S2.xlsx',sheet_name='Table S2')
assert chk1.shape==s1.shape and chk2.shape==s2.shape
num=[c for c in s1.columns if s1[c].dtype.kind=='f']
assert (chk1[num].fillna(-9)-s1[num].fillna(-9)).abs().max().max()<1e-9
print('xlsx ok', chk1.shape, chk2.shape)
